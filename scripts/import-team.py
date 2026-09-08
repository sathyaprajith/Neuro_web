from pathlib import Path
import json
from urllib.parse import urlparse

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = PROJECT_ROOT / "Student_details.xlsx"
OUTPUT_PATH = PROJECT_ROOT / "src" / "data" / "team.ts"


def normalize_url(value):
    text = str(value or "").strip()
    if not text:
        return None
    candidate = text if text.lower().startswith(("http://", "https://")) else f"https://{text}"
    parsed = urlparse(candidate)
    return candidate if parsed.scheme == "https" and parsed.netloc else None


workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
sheet = workbook[workbook.sheetnames[0]]
headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
indices = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
team_members = []

for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
    name = str(row[indices["Name"]] or "").strip()
    if not name:
        continue
    team_members.append(
        {
            "id": f"member-{row_number}",
            "name": name,
            **({"linkedin": normalize_url(row[indices["Linkedin link"]])} if "Linkedin link" in indices and normalize_url(row[indices["Linkedin link"]]) else {}),
            **({"github": normalize_url(row[indices["Github link"]])} if "Github link" in indices and normalize_url(row[indices["Github link"]]) else {}),
            **({"image": normalize_url(row[indices["Photo"]])} if "Photo" in indices and normalize_url(row[indices["Photo"]]) else {}),
        }
    )

content = """export interface TeamMember {
  id: string;
  name: string;
  role?: string;
    bio?: string;
  image?: string;
  linkedin?: string;
  github?: string;
}

export const TEAM_MEMBERS: TeamMember[] = %s;
""" % json.dumps(team_members, indent=2, ensure_ascii=True)
OUTPUT_PATH.write_text(content + "\n", encoding="utf-8")
print(f"Imported {len(team_members)} students into {OUTPUT_PATH.relative_to(PROJECT_ROOT)}")
